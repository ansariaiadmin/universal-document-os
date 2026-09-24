
from fastapi import FastAPI, Request, UploadFile, File, Form
from fastapi.responses import HTMLResponse, JSONResponse, FileResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from pathlib import Path
import uuid, shutil, mimetypes, json, subprocess, time, re

BASE = Path(__file__).resolve().parent.parent
UPLOADS = BASE/"data/uploads"
OUTPUTS = BASE/"data/outputs"
WORKROOMS = BASE/"data/workrooms"
for p in (UPLOADS, OUTPUTS, WORKROOMS): p.mkdir(parents=True, exist_ok=True)

app = FastAPI(title="Universal Document OS", version="1.0.0")
app.mount("/static", StaticFiles(directory=BASE/"app/static"), name="static")
templates = Jinja2Templates(directory=BASE/"app/templates")

def audit(event, **kw):
    rec = {"ts": time.time(), "event": event, **kw}
    with open(BASE/"data/audit.jsonl","a",encoding="utf-8") as f:
        f.write(json.dumps(rec, ensure_ascii=False)+"\n")

def detect(path):
    ext = path.suffix.lower()
    mapping = {
        ".pdf":"PDF", ".docx":"DOCX", ".doc":"DOC", ".xlsx":"XLSX", ".xls":"XLS",
        ".pptx":"PPTX", ".ppt":"PPT", ".odt":"ODT", ".ods":"ODS", ".odp":"ODP",
        ".rtf":"RTF", ".csv":"CSV", ".txt":"TXT", ".md":"MARKDOWN", ".html":"HTML",
        ".htm":"HTML", ".json":"JSON"
    }
    return mapping.get(ext, mimetypes.guess_type(path.name)[0] or "UNKNOWN")

def extract_text(path):
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            return "\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)
        if ext == ".docx":
            from docx import Document
            d=Document(str(path))
            return "\n".join(p.text for p in d.paragraphs)
        if ext == ".xlsx":
            from openpyxl import load_workbook
            wb=load_workbook(path, read_only=True, data_only=False)
            out=[]
            for ws in wb.worksheets:
                out.append(f"[SHEET] {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    out.append("\t".join("" if v is None else str(v) for v in row))
            return "\n".join(out)
        if ext in {".txt",".md",".csv",".json",".html",".htm",".rtf"}:
            return path.read_text(encoding="utf-8", errors="replace")
    except Exception as e:
        return f"[EXTRACTION_ERROR] {e}"
    return ""

@app.get("/", response_class=HTMLResponse)
async def landing(request: Request):
    return templates.TemplateResponse("landing.html", {"request": request})

@app.get("/app", response_class=HTMLResponse)
async def panel(request: Request):
    return templates.TemplateResponse("panel.html", {"request": request})

@app.get("/api/health")
async def health():
    return {"status":"ok","service":"Universal Document OS","version":"1.0.0"}

@app.post("/api/process")
async def process(file: UploadFile=File(...), operation: str=Form("analyze"), target_format: str=Form("same")):
    job = uuid.uuid4().hex
    safe = Path(file.filename or "upload.bin").name
    src = UPLOADS/f"{job}_{safe}"
    with src.open("wb") as out:
        shutil.copyfileobj(file.file, out)
    fmt = detect(src)
    text = extract_text(src)
    result = {
        "job_id":job, "filename":safe, "detected_format":fmt,
        "operation":operation, "target_format":target_format,
        "characters":len(text), "status":"ANALYZED",
        "preview":text[:5000]
    }
    (WORKROOMS/f"{job}.json").write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8")
    audit("ANALYZE", job_id=job, filename=safe, format=fmt, operation=operation)
    # For text-like formats, produce a real downloadable output.
    if operation in {"copy","export_text"} and text and target_format in {"txt","md"}:
        outp=OUTPUTS/f"{job}.{target_format}"
        outp.write_text(text,encoding="utf-8")
        result["download"]=f"/api/download/{outp.name}"
        result["status"]="READY"
    return JSONResponse(result)

@app.get("/api/download/{name}")
async def download(name: str):
    p=OUTPUTS/name
    if not p.exists() or p.parent != OUTPUTS:
        return JSONResponse({"error":"not found"},status_code=404)
    return FileResponse(p, filename=p.name)

@app.get("/api/status")
async def status():
    return {"uploads":len(list(UPLOADS.iterdir())),"outputs":len(list(OUTPUTS.iterdir())),"workrooms":len(list(WORKROOMS.iterdir()))}
